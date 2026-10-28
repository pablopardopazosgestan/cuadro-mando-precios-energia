"""MIBGAS: índice diario MIBGAS-ES (€/MWh) desde los Excel anuales MIBGAS_Data_AAAA.xlsx.

Estructura comprobada en la fase 1:
- Hasta 2022: hoja `Indices`, columnas `Delivery day`, `Area`, `MIBGAS Index [EUR/MWh]` (2022) o
  `MIBGAS-ES Index [EUR/MWh]` (2020–2021), con una fila por área (ES/PT).
- 2023 en adelante: hoja `MIBGAS Indexes`, columna `MIBGAS-ES Index [EUR/MWh]`, una fila por día.
La hoja `Definitions` define igual ambas series (precio medio ponderado de todas las
transacciones con entrega en el mismo día de gas en el PVB).
"""
from __future__ import annotations

import datetime as dt
import io
import logging
from dataclasses import dataclass

import openpyxl

from ..http import clave_cache, descargar

log = logging.getLogger(__name__)

URL = "https://www.mibgas.es/es/file-access/MIBGAS_Data_{anyo}.xlsx?path=AGNO_{anyo}/XLS"
FUENTE = "MIBGAS, fichero anual MIBGAS_Data (hoja de índices), MIBGAS-ES Index"
URL_FUENTE = "https://www.mibgas.es/es/file-access"


@dataclass(frozen=True)
class DiaMibgas:
    fecha: dt.date
    indice_es: float  # €/MWh
    volumen_mwh: float | None


def leer_excel(contenido: bytes) -> list[DiaMibgas]:
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hojas = [h for h in wb.sheetnames if "ndex" in h or "ndice" in h]
    if not hojas:
        raise ValueError(f"el Excel de MIBGAS no tiene hoja de índices: {wb.sheetnames}")
    ws = wb[hojas[0]]
    filas = ws.iter_rows(values_only=True)
    cab = [str(c).replace("\n", " ").strip() if c is not None else "" for c in next(filas)]
    out: list[DiaMibgas] = []
    if "Area" in cab:  # formato hasta 2022: una fila por área; la columna se llama «MIBGAS-ES Index» (2020–21) o «MIBGAS Index» (2022)
        i_area = cab.index("Area")
        i_idx = next(i for i, c in enumerate(cab) if c.startswith("MIBGAS Index") or c.startswith("MIBGAS-ES Index"))
        i_vol = next((i for i, c in enumerate(cab) if c.startswith("MIBGAS Volume") or c.startswith("MIBGAS-ES Volume")), None)
        for f in filas:
            if isinstance(f[0], dt.datetime) and f[i_area] == "ES" and f[i_idx] is not None:
                out.append(DiaMibgas(f[0].date(), float(f[i_idx]), float(f[i_vol]) if i_vol is not None and f[i_vol] is not None else None))
    else:  # formato 2023+
        i_idx = next(i for i, c in enumerate(cab) if c.startswith("MIBGAS-ES Index"))
        for f in filas:
            if isinstance(f[0], dt.datetime) and f[i_idx] is not None:
                out.append(DiaMibgas(f[0].date(), float(f[i_idx]), None))
    out.sort(key=lambda d: d.fecha)
    return out


def descargar_anyo(anyo: int, *, cache: bool) -> list[DiaMibgas]:
    """Descarga el Excel del año. Los años cerrados se cachean; el actual se pide siempre."""
    cuerpo = descargar(URL.format(anyo=anyo), cache=clave_cache("mibgas", f"MIBGAS_Data_{anyo}.xlsx") if cache else None)
    return leer_excel(cuerpo)


def descargar_rango(desde: dt.date, hasta: dt.date) -> list[DiaMibgas]:
    dias: list[DiaMibgas] = []
    hoy = dt.date.today()
    for anyo in range(desde.year, hasta.year + 1):
        dias += [d for d in descargar_anyo(anyo, cache=anyo < hoy.year) if desde <= d.fecha <= hasta]
    return dias
