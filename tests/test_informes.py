"""Generación de informes: prueba de humo sobre los datos reales del repositorio."""
import datetime as dt
from pathlib import Path

import openpyxl
import pdfplumber
import pytest

from precios import informes

DATOS = Path(__file__).resolve().parent.parent / "datos" / "series"


@pytest.mark.skipif(not (DATOS / "omie.csv").exists(), reason="sin datos")
def test_generar_pdf_y_excel(tmp_path, monkeypatch):
    monkeypatch.setattr(informes, "DIR_INFORMES", tmp_path)
    salidas = informes.generar(hoy=dt.date(2026, 10, 27))
    pdf, xlsx = salidas
    assert pdf.exists() and xlsx.exists()
    with pdfplumber.open(pdf) as doc:
        assert len(doc.pages) == 2
        texto = doc.pages[0].extract_text()
        assert "Claves de la semana" in texto and "OMIE" in texto and "€/kWh" in texto
        assert "Pellet frente a las alternativas" in texto
        assert "Qué significa cada indicador" in doc.pages[1].extract_text()
    wb = openpyxl.load_workbook(xlsx)
    assert {"Cuadro de mando", "Anual", "Mensual", "Gráficos", "Fuentes y criterios"} <= set(wb.sheetnames)
    assert len(wb["Gráficos"]._charts) == 11            # comparativo + barras anuales + 9 individuales
    assert len(wb["Cuadro de mando"]._charts) == 1
    datos_omie = wb["Datos 1 OMIE"]
    assert "Datos_omie" in datos_omie.tables              # tabla de Excel filtrable
    assert datos_omie.max_row > 2000                      # histórico diario completo
    assert wb["Mensual"]["A2"].value == "2020-01"
    assert wb["Anual"]["A2"].value == 2020


def test_etiquetas_finales_separadas_y_dentro_del_eje():
    import matplotlib.pyplot as plt
    fig, ax = plt.subplots(figsize=(6, 3), dpi=100)
    ax.set_ylim(0, 1)
    ax.set_xlim(0, 10)
    items = [("tur2", 0.50, "TUR 2 0,500"), ("butano", 0.505, "Butano 0,505"), ("gasoleo_c", 0.51, "Gasóleo C 0,510"), ("pvpc", 0.995, "PVPC 0,995")]
    ys = informes._etiquetas_finales(ax, items, 10.2, fontsize=7)
    alto_px = ax.get_window_extent().height
    px = sorted((y - 0) / 1 * alto_px for y in ys)
    minimo = 7 * 1.5 * fig.dpi / 72
    assert len(ys) == 4
    assert all(b - a >= minimo - 1e-6 for a, b in zip(px, px[1:]))     # ninguna etiqueta pisa a otra
    assert max(ys) <= 1.0 + 1e-9                                        # la de arriba no se sale del eje
    assert len(ax.texts) == 4 and ax.texts[0].get_position()[0] == 10.2  # a la derecha del último dato
    plt.close(fig)
    assert informes._etiquetas_finales(ax, [], 1) == []


@pytest.mark.skipif(not (DATOS / "omie.csv").exists(), reason="sin datos")
def test_comparativo_energia_final_y_periodo_sombreado():
    png = informes.grafico_comparativo(dt.date(2022, 1, 1), dt.date(2026, 10, 27), ids=["tur2", "gasoleo_c"], alto=3.4)
    assert png[:8] == b"\x89PNG\r\n\x1a\n"
    png2 = informes.grafico_individual("tur2", dt.date(2022, 1, 1), dt.date(2026, 10, 27), (dt.date(2026, 10, 21), dt.date(2026, 10, 27)))
    assert png2[:8] == b"\x89PNG\r\n\x1a\n"


@pytest.mark.skipif(not (DATOS / "omie.csv").exists(), reason="sin datos")
def test_pdf_dos_paginas_y_excel_comparador(tmp_path, monkeypatch):
    monkeypatch.setattr(informes, "DIR_INFORMES", tmp_path)
    pdf, xlsx = informes.generar(hoy=dt.date(2026, 10, 27))
    with pdfplumber.open(pdf) as doc:
        assert len(doc.pages) == 2
        p1, p2 = doc.pages[0].extract_text(), doc.pages[1].extract_text()
        assert "Sobrecoste" in p1 and len(doc.pages[0].images) >= 10            # 8 mini tendencias + ranking + competitividad
        for nombre in ("1. OMIE", "2. MIBGAS", "3. Gasóleo A", "4. PVPC 2.0TD", "5. TUR 2", "6. Gasóleo C", "7. Butano", "8. Pellet saco 15 kg"):
            assert nombre in p1 and nombre in p2                                 # los ocho indicadores en las dos páginas
        assert "Granel" in p1 and "Granel" in p2
        assert "Qué mide y por qué importa" in p2 and "Esta semana (calculado)" in p2 and "24 meses" in p2
        assert "Fuentes en la última actualización" in p2 and "Cómo se calcula" in p2
    wb = openpyxl.load_workbook(xlsx)
    assert wb.sheetnames.index("Comparador") < wb.sheetnames.index("Datos 1 OMIE")
    wc = wb["Comparador"]
    assert len(wc.data_validations.dataValidation) == 3 and len(wc._charts) == 1
    assert wc["C5"].value == "Gasóleo C" and wc["C9"].value == "Mensual" and wc["C10"].value == "No"
    assert wc["B14"].value.startswith("=IF($C$5=") and "MATCH" in wc["H12"].value and "INDEX(Mensual!" in wc["H15"].value
    assert wc["B15"].value == '=IF($C$10="Sí",H15/H$13*100,H15)'
    assert wc["P5"].value == "OMIE" and wc["Q5"].value == wb["Mensual"]["B1"].value and wc["R5"].value == wb["Anual"]["B1"].value
    ws = wb["Cuadro de mando"]
    assert len(ws._charts) == 1 and len(ws.merged_cells.ranges) >= 30            # tarjetas en celdas combinadas
    assert ws["C5"].value.startswith("1. OMIE") and isinstance(ws["C6"].value, float)
    fila = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=2).value or "").startswith("Ranking"))
    precios = [ws.cell(row=r, column=4).value for r in range(fila + 2, fila + 9)]
    assert precios == sorted(precios) and len(precios) == 7                      # energía final, de menor a mayor
    assert wb.calculation.fullCalcOnLoad is True
